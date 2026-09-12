from pathlib import Path
import pymupdf as fitz
from pptx import Presentation

def parse_pdf(file_path: str) -> list[dict]:
    """Extracts text from a PDF file page by page."""
    doc = fitz.open(file_path)
    extracted_data = []
    
    for page_num in range(len(doc)):
        page = doc[page_num]
        text = page.get_text("text").strip()
        if text:
            extracted_data.append({
                "source": Path(file_path).name,
                "page_number": page_num + 1,
                "text": text
            })
    return extracted_data

def parse_pptx(file_path: str) -> list[dict]:
    """Extracts text from a PowerPoint file slide by slide."""
    prs = Presentation(file_path)
    extracted_data = []
    
    for slide_num, slide in enumerate(prs.slides, start=1):
        slide_text = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        slide_text.append(text)
        
        full_text = "\n".join(slide_text)
        if full_text:
            extracted_data.append({
                "source": Path(file_path).name,
                "page_number": slide_num,
                "text": full_text
            })
    return extracted_data

def parse_markdown(file_path: str) -> list[dict]:
    """Reads text from a Markdown (.md) or plain text file."""
    with open(file_path, "r", encoding="utf-8") as f:
        text = f.read().strip()
        
    if not text:
        return []
        
    return [{
        "source": Path(file_path).name,
        "page_number": 1,
        "text": text
    }]

def parse_document(file_path: str) -> list[dict]:
    """Main router function to parse supported files based on extension."""
    path = Path(file_path)
    ext = path.suffix.lower()
    
    if ext == ".pdf":
        return parse_pdf(file_path)
    elif ext == ".pptx":
        return parse_pptx(file_path)
    elif ext in [".md", ".txt"]:
        return parse_markdown(file_path)
    else:
        raise ValueError(f"Unsupported file format: {ext}")